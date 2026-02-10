"""Document API routes."""
from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import Response

from api.schemas.document import DocumentResponse, DocumentCreate
from application.services.document_service import DocumentService

router = APIRouter()
document_service = DocumentService()


@router.post("/documents", response_model=DocumentResponse, status_code=201)
async def upload_document(file: UploadFile = File(...)):
    """Upload a document file.
    
    Args:
        file: Uploaded file
        
    Returns:
        DocumentResponse with document information
    """
    try:
        document = await document_service.create_from_upload(file)
        return DocumentResponse.from_domain(document)
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to upload document: {str(e)}"
        )


@router.post("/documents/url", response_model=DocumentResponse, status_code=201)
async def upload_document_from_url(request: DocumentCreate):
    """Upload a document from URL.
    
    Args:
        request: Document creation request with URL
        
    Returns:
        DocumentResponse with document information
        
    Note:
        Future feature - not implemented yet
    """
    raise HTTPException(
        status_code=501,
        detail="URL upload not implemented yet"
    )


def _get_xlsx_workbook(document):
    """Validate document is xlsx and return open workbook. Caller must close."""
    from pathlib import Path
    from openpyxl import load_workbook
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    ft = (document.file_type or "").strip().lower()
    if ft not in ("xlsx", "xls"):
        raise HTTPException(
            status_code=404,
            detail=f"Preview only supported for xlsx (document type is '{document.file_type}')",
        )
    if not Path(document.file_path).is_file():
        raise HTTPException(status_code=404, detail="File not found on server")
    return load_workbook(document.file_path, read_only=True, data_only=True)


@router.get("/documents/{document_id}/xlsx-sheets")
async def get_document_xlsx_sheets(document_id: str):
    """Get list of sheet names for xlsx. Returns 404 for non-xlsx."""
    document = document_service.get_by_id(document_id)
    wb = _get_xlsx_workbook(document)
    try:
        sheets = [{"index": i, "name": ws.title} for i, ws in enumerate(wb.worksheets)]
        return {"sheets": sheets}
    finally:
        wb.close()


@router.get("/documents/{document_id}/xlsx-preview")
async def get_document_preview(document_id: str, sheet: int = 0):
    """Get xlsx preview for one sheet (0-based index). Returns 404 for non-xlsx."""
    document = document_service.get_by_id(document_id)
    wb = _get_xlsx_workbook(document)
    try:
        if sheet < 0 or sheet >= len(wb.worksheets):
            raise HTTPException(
                status_code=400,
                detail=f"sheet index must be 0..{len(wb.worksheets) - 1}",
            )
        ws = wb.worksheets[sheet]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else "" for c in row])
        return {"sheet_index": sheet, "sheet_name": ws.title, "rows": rows}
    finally:
        wb.close()


@router.get("/documents/{document_id}", response_model=DocumentResponse)
async def get_document(document_id: str):
    """Get document information.
    
    Args:
        document_id: Document ID
        
    Returns:
        Document response with document information
    """
    document = document_service.get_by_id(document_id)
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    return DocumentResponse.from_domain(document)


@router.get("/documents/{document_id}/pages/{page_number}")
async def get_page_image(document_id: str, page_number: int):
    """Get page image for document viewer.
    
    Args:
        document_id: Document ID
        page_number: Page number (0-indexed)
        
    Returns:
        Image bytes (PNG/JPEG)
        
    Note:
        Supports page navigation (Prev, Next)
    """
    # Get document
    document = document_service.get_by_id(document_id)
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    # Validate page number
    if page_number < 0:
        raise HTTPException(status_code=400, detail="Page number must be >= 0")
    if page_number >= document.page_count:
        raise HTTPException(
            status_code=404,
            detail=f"Page {page_number} not found. Document has {document.page_count} page(s)"
        )
    
    # Get page image
    image_bytes = document_service.get_page_image(document, page_number)
    if not image_bytes:
        raise HTTPException(status_code=404, detail="Page image not found")
    
    # Determine content type
    content_type = "image/png"
    if document.is_image():
        if document.file_type.lower() in ["jpg", "jpeg"]:
            content_type = "image/jpeg"
        elif document.file_type.lower() == "webp":
            content_type = "image/webp"
    
    return Response(content=image_bytes, media_type=content_type)


@router.delete("/documents/{document_id}", status_code=204)
async def delete_document(document_id: str):
    """Delete a document.
    
    Args:
        document_id: Document ID
        
    Note:
        Future feature - not implemented yet
    """
    raise HTTPException(
        status_code=501,
        detail="Document deletion not implemented yet"
    )
