#!/usr/bin/env python3
"""E2E Chunking test via Python. Run from playground root.
Prerequisites: Redis, Peter-parser (8001), Peter-parser RQ worker, Playground backend (8000).
"""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

try:
    import httpx
except ImportError:
    print("pip/uv install httpx first")
    sys.exit(1)

# Allow override for CI or remote backend
API_BASE = os.environ.get("API_BASE", "http://localhost:8000/api/v1").rstrip("/")
if not API_BASE.endswith("/api/v1"):
    API_BASE = f"{API_BASE.rstrip('/')}/api/v1"
HEALTH_URL = API_BASE.replace("/api/v1", "") + "/health"

SAMPLE = Path(__file__).resolve().parent.parent / "sample_excel.xlsx"
TIMEOUT_POST = 100.0  # backend may wait for Peter-parser up to 90s
TIMEOUT_GET = 10.0
TIMEOUT_HEALTH = 3.0
POLL_INTERVAL = 2
POLL_ATTEMPTS = 60


def ensure_sample_excel() -> Path:
    """Create minimal sample_excel.xlsx if missing."""
    if SAMPLE.is_file():
        return SAMPLE
    try:
        from openpyxl import Workbook
    except ImportError:
        print(f"Error: {SAMPLE} not found. Create it or add openpyxl and re-run.")
        sys.exit(1)
    wb = Workbook()
    ws = wb.active
    if ws:
        ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "제품명", "수량", "단가", "합계"
        ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "노트북", 2, 1500000, 3000000
    SAMPLE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAMPLE)
    print(f"Created {SAMPLE}")
    return SAMPLE


def check_backend_health() -> bool:
    """Return True if backend is up."""
    try:
        r = httpx.get(HEALTH_URL, timeout=TIMEOUT_HEALTH)
        return r.status_code == 200
    except Exception:
        return False


def main() -> int:
    sample_path = ensure_sample_excel()

    print("=== E2E Chunking Test (Python) ===")
    print(f"  Backend health: {HEALTH_URL}")
    if not check_backend_health():
        print("  FAIL: Backend is not running. Start:")
        print("    1. Redis (6379)")
        print("    2. Peter-parser: cd ../pipelines/parsing-pipeline/peter-parser && uv run python main.py  (from playground)")
        print("    3. Peter-parser RQ worker: cd ../pipelines/parsing-pipeline/peter-parser && uv run python -m peter_parser.worker")
        print("    4. Backend: cd backend && uv run python main.py  (from playground)")
        return 1
    print("  OK")

    print(f"1. POST {API_BASE}/chunking/parse ...")
    try:
        with open(sample_path, "rb") as f:
            content = f.read()
        r = httpx.post(
            f"{API_BASE}/chunking/parse",
            files={"file": (sample_path.name, content)},
            data={"document_type": "excel"},
            timeout=TIMEOUT_POST,
        )
    except httpx.ConnectError as e:
        print(f"   Connection failed: {e}")
        print("   Ensure Playground backend is running on 8000 (or set API_BASE).")
        return 1
    except httpx.TimeoutException as e:
        print(f"   Timeout after {TIMEOUT_POST}s.")
        print("   Backend may be waiting for Peter-parser. Is Peter-parser (8001) running?")
        return 1

    if r.status_code == 502:
        print(f"   HTTP 502: {r.text[:400]}")
        print("   Is Peter-parser (port 8001) running? Is Redis running?")
        return 1
    if r.status_code != 202:
        print(f"   HTTP {r.status_code}: {r.text[:300]}")
        return 1

    data = r.json()
    job_id = data.get("job_id") or ""
    doc_id = data.get("document_id") or ""
    if not job_id:
        print("   No job_id in response:", data)
        return 1
    print(f"   job_id={job_id} document_id={doc_id}")

    print(f"2. Polling GET /chunking/status/{job_id} ...")
    for i in range(POLL_ATTEMPTS):
        try:
            r2 = httpx.get(f"{API_BASE}/chunking/status/{job_id}", timeout=TIMEOUT_GET)
        except (httpx.ConnectError, httpx.TimeoutException) as e:
            print(f"   Attempt {i+1}: request failed: {e}")
            time.sleep(POLL_INTERVAL)
            continue
        st = r2.json().get("status", "")
        print(f"   Attempt {i+1}: status={st}")
        if st == "completed":
            break
        if st == "failed":
            err = r2.json().get("error", "")
            print(f"   Job failed: {err}")
            return 1
        time.sleep(POLL_INTERVAL)
    else:
        print("   Timeout waiting for completed. Is Peter-parser RQ worker running?")
        return 1

    print(f"3. GET /chunking/result/{job_id} ...")
    try:
        r3 = httpx.get(f"{API_BASE}/chunking/result/{job_id}", timeout=TIMEOUT_GET)
    except (httpx.ConnectError, httpx.TimeoutException) as e:
        print(f"   Failed: {e}")
        return 1
    if r3.status_code != 200:
        print(f"   HTTP {r3.status_code}: {r3.text[:200]}")
        return 1
    result = r3.json()
    chunks = result.get("chunks") or []
    print(f"   Chunks: {len(chunks)}")
    for i, c in enumerate(chunks[:3]):
        text = (c.get("chunk") or "")[:80]
        print(f"   - #{i+1}: {text}...")

    print("=== E2E Test Complete ===")
    return 0


if __name__ == "__main__":
    sys.exit(main())
